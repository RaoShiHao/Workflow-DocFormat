"""Fixed-seed assignment of dataset factors per topic×complexity.

Nested factor layout (see ``sample_info.json`` v6):

``init`` (shared per content_key — hierarchy aids on ``input.docx`` / golden):

  - ``heading_styles`` — apply built-in heading styles on init (partial vs plain)
  - ``numbering.enabled`` + ``numbering.scheme_id`` — text-prefix numbering on titles
    (if false, no scheme draw and no downstream numbering choices)

``text`` (only when ``input_mode`` uses ``text_input.md``):

  - ``target_notes`` — include a one-line purpose per style/section (1:1 draw)
  - ``numbering_note`` — per-heading numbering hint in text (1:1; only when numbering on)

``input_mode`` implies which files exist (template / text); no redundant booleans.

Input modes (exactly **one** per topic×Cx; modes are balanced across content units):

  - ``text_requirement`` — text-only format spec
  - ``template`` — ``template.docx`` only
  - ``template_w_text`` — template + text (catalog / delta / both at build time)
"""
from __future__ import annotations

import hashlib
import random
from pathlib import Path
from typing import Any

import yaml

from LongDocFmtEval.common.themes_xlsx import COMPLEXITY_TIERS
from LongDocFmtEval.pipeline_auto.auto_paths import CONFIG_DIR

DEFAULT_POOLS_PATH = CONFIG_DIR / "training_factor_pools.yaml"
DEFAULT_GENRES_XLSX = CONFIG_DIR / "genres.xlsx"
# Frozen full-set assignment (eval selection by topic_id elsewhere; no train split).
FROZEN_SEED = 1234
DEFAULT_OUTPUT = CONFIG_DIR / "datasets_factors.json"
LEGACY_OUTPUT = CONFIG_DIR / "training_factors.json"

INIT_STRUCTURES = ("plain", "partial")
INPUT_MODES = ("text_requirement", "template", "template_w_text")

_LEGACY_FLAT_KEYS = (
    "numbering_enabled",
    "scheme_id",
    "init_structure",
    "document_numbered",
    "init_text_numbered",
    "template_numbered",
    "text_numbering_spec",
    "text_target_notes",
    "uses_template_docx",
    "uses_text_input",
    "text_input_content",
    "template_mode",
)


def rng_for(key: str, master_seed: int = 42) -> random.Random:
    digest = hashlib.sha256(f"{master_seed}:{key}".encode("utf-8")).hexdigest()
    return random.Random(int(digest[:16], 16))


def _weighted_choice(rng: random.Random, items: list[str], weights: list[float]) -> str:
    if not items:
        raise ValueError("empty choice pool")
    if len(weights) != len(items):
        weights = [1.0] * len(items)
    return rng.choices(items, weights=weights, k=1)[0]


def load_pools(path: Path | None = None) -> dict[str, Any]:
    p = path or DEFAULT_POOLS_PATH
    raw = yaml.safe_load(p.read_text(encoding="utf-8")) if p.is_file() else {}
    return raw if isinstance(raw, dict) else {}


def _bool_choice(rng: random.Random, p_true: float) -> bool:
    return rng.random() < float(p_true)


def init_block(factors: dict[str, Any]) -> dict[str, Any]:
    """Canonical ``init`` view (nested or legacy flat assignment row)."""
    if isinstance(factors.get("init"), dict):
        blk = factors["init"]
        num = blk.get("numbering") if isinstance(blk.get("numbering"), dict) else {}
        enabled = bool(num.get("enabled"))
        scheme = num.get("scheme_id") if enabled else None
        return {
            "heading_styles": bool(blk.get("heading_styles")),
            "numbering": {"enabled": enabled, "scheme_id": scheme},
        }
    num = bool(factors.get("numbering_enabled"))
    return {
        "heading_styles": str(factors.get("init_structure") or "plain") == "partial",
        "numbering": {
            "enabled": num,
            "scheme_id": factors.get("scheme_id") if num else None,
        },
    }


def init_heading_styles(factors: dict[str, Any]) -> bool:
    return bool(init_block(factors).get("heading_styles"))


def init_structure_name(factors: dict[str, Any]) -> str:
    return "partial" if init_heading_styles(factors) else "plain"


def numbering_enabled(factors: dict[str, Any]) -> bool:
    return bool(init_block(factors)["numbering"]["enabled"])


def scheme_id(factors: dict[str, Any]) -> str | None:
    num = init_block(factors)["numbering"]
    if not num.get("enabled"):
        return None
    sid = num.get("scheme_id")
    return str(sid) if sid else None


def uses_template_docx(factors: dict[str, Any]) -> bool:
    return str(factors.get("input_mode") or "") in ("template", "template_w_text")


def uses_text_input(factors: dict[str, Any]) -> bool:
    return str(factors.get("input_mode") or "") in ("text_requirement", "template_w_text")


def text_block(factors: dict[str, Any]) -> dict[str, Any]:
    tb = factors.get("text")
    if isinstance(tb, dict):
        return dict(tb)
    out: dict[str, Any] = {}
    if uses_text_input(factors):
        out["target_notes"] = bool(factors.get("text_target_notes", True))
        if numbering_enabled(factors):
            out["numbering_note"] = bool(factors.get("text_numbering_spec"))
    return out


def text_numbering_note(factors: dict[str, Any]) -> bool:
    if not numbering_enabled(factors) or not uses_text_input(factors):
        return False
    return bool(text_block(factors).get("numbering_note"))


def text_target_notes(factors: dict[str, Any]) -> bool:
    if not uses_text_input(factors):
        return False
    return bool(text_block(factors).get("target_notes", True))


def template_numbered(factors: dict[str, Any]) -> bool:
    return numbering_enabled(factors) and uses_template_docx(factors)


def scheme_pool_for_domain(
    pools: dict[str, Any],
    domain_id: str,
    locale: str,
) -> tuple[list[str], list[float]]:
    loc = "en" if str(locale).lower().startswith("en") else "zh"
    dom = str(domain_id or "").strip()
    block = (pools.get("domain_pools") or {}).get(dom) or pools.get("default_pool") or {}
    entry = block.get(loc) or block.get("zh") or []
    if isinstance(entry, dict):
        items = list(entry.get("pool") or [])
        weights = list(entry.get("weights") or [1.0] * len(items))
    else:
        items = [str(x).strip() for x in entry if str(x).strip()]
        weights = [1.0] * len(items)
    if not items:
        items = ["zh_section_decimal" if loc == "zh" else "en_chapter_decimal"]
        weights = [1.0]
    return items, weights


def assign_content_factors(
    topic_id: int,
    complexity: str,
    *,
    pools: dict[str, Any],
    domain_id: str,
    locale: str,
    master_seed: int,
) -> dict[str, Any]:
    """Shared factors for one topic×complexity (2×2 init hierarchy aids + numbering)."""
    key = f"topic:{topic_id}:{complexity}"
    rng = rng_for(key, master_seed)
    inits = list(pools.get("init_structures") or INIT_STRUCTURES)
    i_w = list(pools.get("init_weights") or [0.5, 0.5])
    p_num = float(
        pools.get("numbering_weight", pools.get("document_numbered_weight", 0.5))
    )
    # 2×2 aids: heading styles × text-prefix numbering on titles
    init_structure = _weighted_choice(rng, inits, i_w)
    numbering_enabled = _bool_choice(rng, p_num)
    scheme_id: str | None = None
    if numbering_enabled:
        scheme_rng = rng_for(f"{key}:scheme", master_seed)
        items, weights = scheme_pool_for_domain(pools, domain_id, locale)
        scheme_id = _weighted_choice(scheme_rng, items, weights)
    return {
        "init": {
            "heading_styles": init_structure == "partial",
            "numbering": {
                "enabled": numbering_enabled,
                "scheme_id": scheme_id,
            },
        },
    }


def _draw_text_fields(
    *,
    sample_key: str,
    numbering_on: bool,
    master_seed: int,
    pools: dict[str, Any] | None = None,
) -> dict[str, Any]:
    pools = pools or {}
    p_target = float(pools.get("target_notes_weight", 0.5))
    p_num_note = float(pools.get("text_numbering_note_weight", 0.5))
    sk = sample_key or "_"
    text: dict[str, Any] = {
        "target_notes": _bool_choice(rng_for(f"{sk}:target_notes", master_seed), p_target),
    }
    if numbering_on:
        text["numbering_note"] = _bool_choice(
            rng_for(f"{sk}:text_num_spec", master_seed), p_num_note,
        )
    return text


def derive_mode_fields(
    input_mode: str,
    base: dict[str, Any],
    *,
    sample_key: str = "",
    master_seed: int = FROZEN_SEED,
    pools: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Map shared content ``init`` → per-mode row (adds ``text`` when applicable)."""
    mode = str(input_mode or "").strip().lower()
    if mode not in INPUT_MODES:
        raise ValueError(f"unknown input_mode: {input_mode!r}")
    init = init_block({**base, "init": base.get("init")})
    row: dict[str, Any] = {"input_mode": mode, "init": init}
    if uses_text_input(row):
        row["text"] = _draw_text_fields(
            sample_key=sample_key or f"_:{mode}",
            numbering_on=bool(init["numbering"]["enabled"]),
            master_seed=master_seed,
            pools=pools,
        )
    return row


def resolve_runtime_factors(
    factors: dict[str, Any],
    *,
    master_seed: int = FROZEN_SEED,
    pools: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Normalize to nested ``init`` / ``text``; draw missing text flags at build time."""
    out = dict(factors)
    sample_key = str(out.get("sample_key") or "")
    init = init_block(out)
    if not init["numbering"]["enabled"]:
        init["numbering"]["scheme_id"] = None
    out["init"] = init

    if uses_text_input(out):
        text = text_block(out)
        if "target_notes" not in text:
            drawn = _draw_text_fields(
                sample_key=sample_key,
                numbering_on=bool(init["numbering"]["enabled"]),
                master_seed=master_seed,
                pools=pools,
            )
            text.setdefault("target_notes", drawn["target_notes"])
            if init["numbering"]["enabled"]:
                text.setdefault("numbering_note", drawn.get("numbering_note", False))
        if not init["numbering"]["enabled"]:
            text.pop("numbering_note", None)
        out["text"] = text
    else:
        out.pop("text", None)

    for key in _LEGACY_FLAT_KEYS:
        out.pop(key, None)
    return out


def validate_domain_pools(
    pools: dict[str, Any],
    *,
    known_schemes: set[str] | None = None,
    expected_per_locale: int = 3,
) -> list[str]:
    errors: list[str] = []
    known = known_schemes or set()
    for dom, block in (pools.get("domain_pools") or {}).items():
        if not isinstance(block, dict):
            errors.append(f"domain_pools.{dom}: expected mapping of locale → scheme list")
            continue
        for loc, entry in block.items():
            if isinstance(entry, dict):
                items = list(entry.get("pool") or [])
            else:
                items = [str(x).strip() for x in entry if str(x).strip()]
            if len(items) != expected_per_locale:
                errors.append(
                    f"domain_pools.{dom}.{loc}: expected {expected_per_locale} schemes, got {len(items)}"
                )
            if len(set(items)) != len(items):
                errors.append(f"domain_pools.{dom}.{loc}: duplicate scheme ids {items}")
            for sid in items:
                if known and sid not in known:
                    errors.append(f"domain_pools.{dom}.{loc}: unknown scheme_id {sid!r}")
    return errors


def assign_topic_scheme(
    topic_id: int,
    *,
    domain_id: str,
    locale: str,
    pools: dict[str, Any],
    master_seed: int,
) -> str:
    rng = rng_for(f"topic:{topic_id}:scheme", master_seed)
    items, weights = scheme_pool_for_domain(pools, domain_id, locale)
    return _weighted_choice(rng, items, weights)


# Back-compat alias
assign_sample_factors = assign_content_factors


def _read_xlsx_rows(path: Path) -> tuple[list[str], list[tuple]]:
    try:
        from openpyxl import load_workbook  # noqa: PLC0415
    except ImportError:
        return [], []
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    if ws.title != "genres" and "genres" in wb.sheetnames:
        ws = wb["genres"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(h or "").strip().lower() for h in rows[0]]
    return headers, rows[1:]


def load_genre_rows(xlsx: Path | None = None) -> list[dict[str, Any]]:
    path = xlsx or DEFAULT_GENRES_XLSX
    headers, body = _read_xlsx_rows(path)
    if not headers:
        return []
    out: list[dict[str, Any]] = []
    for raw in body:
        if not raw:
            continue
        item = {headers[i]: raw[i] for i in range(min(len(headers), len(raw)))}
        try:
            gid = int(item.get("id") or 0)
        except (TypeError, ValueError):
            continue
        enabled = str(item.get("enabled") or "1").strip()
        if enabled not in ("1", "true", "yes", "Y", "y"):
            continue
        locale = str(item.get("locale") or "zh").strip() or "zh"
        coarse = str(item.get("domain_id") or "").strip()
        fine = str(item.get("fine_domain_id") or coarse or "").strip()
        out.append({
            "topic_id": gid,
            "domain_id": coarse,
            "fine_domain_id": fine,
            "domain": str(item.get("domain") or "").strip(),
            "locale": "en" if locale.lower().startswith("en") else "zh",
            "genre": str(item.get("genre") or "").strip(),
            "genre_display_name": str(item.get("genre_display_name") or "").strip(),
        })
    return out


def build_assignments(
    genre_rows: list[dict[str, Any]] | None = None,
    *,
    pools: dict[str, Any] | None = None,
    master_seed: int = FROZEN_SEED,
    complexities: tuple[str, ...] = COMPLEXITY_TIERS,
    input_modes: tuple[str, ...] = INPUT_MODES,
    modes_per_content: int = 1,
) -> dict[str, Any]:
    """Build frozen factor rows.

    ``modes_per_content=1`` (default): each topic×Cx gets **one** input_mode.
    Modes are round-robin over sorted content units so the three modes stay
    roughly equal. Pass ``modes_per_content=3`` only for legacy full Cartesian.
    """
    pools = pools or load_pools()
    meta = pools.get("meta") if isinstance(pools.get("meta"), dict) else {}
    seed = int(master_seed)
    rows = genre_rows if genre_rows is not None else load_genre_rows()
    from LongDocFmtEval.pipeline_auto.seed.heading_numbering import (  # noqa: PLC0415
        load_schemes,
        validate_schemes_automation,
    )

    schemes = load_schemes()
    pool_errors = validate_domain_pools(pools, known_schemes=set(schemes))
    scheme_errors = validate_schemes_automation(schemes)
    if pool_errors or scheme_errors:
        msg = "\n".join(pool_errors + scheme_errors)
        raise ValueError(f"training factor pool / scheme validation failed:\n{msg}")

    mode_list = [str(m).strip().lower() for m in input_modes if str(m).strip()]
    if not mode_list:
        raise ValueError("input_modes must be non-empty")
    for m in mode_list:
        if m not in INPUT_MODES:
            raise ValueError(f"unknown input_mode: {m!r}")
    n_modes = max(1, min(int(modes_per_content), len(mode_list)))

    assignments: list[dict[str, Any]] = []
    topic_ids: set[int] = set()
    content_units = 0
    unit_i = 0
    for row in sorted(rows, key=lambda r: int(r["topic_id"])):
        tid = int(row["topic_id"])
        topic_ids.add(tid)
        for cx in complexities:
            base = assign_content_factors(
                tid,
                cx,
                pools=pools,
                domain_id=str(row.get("domain_id") or ""),
                locale=str(row.get("locale") or "zh"),
                master_seed=seed,
            )
            content_units += 1
            if n_modes >= len(mode_list):
                chosen = list(mode_list)
            elif n_modes == 1:
                chosen = [mode_list[unit_i % len(mode_list)]]
            else:
                start = unit_i % len(mode_list)
                chosen = [mode_list[(start + k) % len(mode_list)] for k in range(n_modes)]
            unit_i += 1
            for im in chosen:
                sample_key = f"{tid}_{cx}_{im}"
                derived = derive_mode_fields(
                    im, base, sample_key=sample_key, master_seed=seed, pools=pools,
                )
                assignments.append({
                    "topic_id": tid,
                    "complexity": cx,
                    "input_mode": im,
                    "sample_key": sample_key,
                    "content_key": f"{tid}_{cx}",
                    "locale": row.get("locale"),
                    "domain_id": row.get("domain_id"),
                    "fine_domain_id": row.get("fine_domain_id"),
                    "genre_display_name": row.get("genre_display_name"),
                    **derived,
                })
    return {
        "meta": {
            "seed": seed,
            "frozen": True,
            "source_of_truth": DEFAULT_OUTPUT.name,
            "version": int(meta.get("version") or 1),
            "n_topics": len(topic_ids),
            "n_content_units": content_units,
            "n_samples": len(assignments),
            "modes_per_content": n_modes,
            "input_modes": list(mode_list),
            "pools_file": str(DEFAULT_POOLS_PATH.name),
        },
        "assignments": assignments,
    }


def assignment_for(
    data: dict[str, Any],
    topic_id: int,
    complexity: str,
    input_mode: str | None = None,
) -> dict[str, Any] | None:
    cx = str(complexity or "").strip().upper()
    if not cx.startswith("C"):
        cx = f"C{int(complexity)}"
    mode = str(input_mode or "").strip().lower() or None
    matches: list[dict[str, Any]] = []
    for row in data.get("assignments") or []:
        if int(row.get("topic_id") or 0) != int(topic_id):
            continue
        if str(row.get("complexity")) != cx:
            continue
        if mode and str(row.get("input_mode") or "") != mode:
            continue
        matches.append(row)
    if not matches:
        return None
    return matches[0]


def assignments_for_content(
    data: dict[str, Any],
    topic_id: int,
    complexity: str,
) -> list[dict[str, Any]]:
    cx = str(complexity or "").strip().upper()
    if not cx.startswith("C"):
        cx = f"C{int(complexity)}"
    return [
        r for r in (data.get("assignments") or [])
        if int(r.get("topic_id") or 0) == int(topic_id) and str(r.get("complexity")) == cx
    ]


def summarize_assignments(data: dict[str, Any]) -> str:
    from collections import Counter

    rows = data.get("assignments") or []
    meta = data.get("meta") or {}
    lines = [
        f"seed={meta.get('seed')} "
        f"topics={meta.get('n_topics')} "
        f"content_units={meta.get('n_content_units')} "
        f"samples={meta.get('n_samples')} "
        f"modes_per_content={meta.get('modes_per_content', '?')}",
        f"input_modes={meta.get('input_modes')}",
        "",
    ]
    for name, extractor in (
        ("input_mode", lambda r: r.get("input_mode")),
        ("init.heading_styles", lambda r: init_block(r).get("heading_styles")),
        ("init.numbering.enabled", lambda r: init_block(r)["numbering"]["enabled"]),
        ("text.target_notes", lambda r: text_block(r).get("target_notes") if uses_text_input(r) else "—"),
        ("text.numbering_note", lambda r: text_block(r).get("numbering_note") if uses_text_input(r) else "—"),
        ("scheme_id", lambda r: scheme_id(r) or "—"),
    ):
        c = Counter("—" if extractor(r) in (None, "") else str(extractor(r)) for r in rows)
        lines.append(f"--- {name} ---")
        for k, n in sorted(c.items()):
            lines.append(f"  {k}: {n}")
        lines.append("")
    by_content: dict[str, list[dict]] = {}
    for r in rows:
        by_content.setdefault(str(r.get("content_key")), []).append(r)
    inconsistent = 0
    multi_mode = 0
    for _ck, group in by_content.items():
        if len(group) > 1:
            multi_mode += 1
        nums = {numbering_enabled(g) for g in group}
        inits = {init_heading_styles(g) for g in group}
        if len(nums) > 1 or len(inits) > 1:
            inconsistent += 1
    lines.append("--- content_key consistency ---")
    lines.append(f"  inconsistent groups: {inconsistent} (expect 0)")
    lines.append(f"  content_keys with >1 mode: {multi_mode}")
    return "\n".join(lines)
