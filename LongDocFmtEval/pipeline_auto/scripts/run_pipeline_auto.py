"""pipeline_auto · 全自动批跑（无 seed / 无人改 Word）。

目录：datasets/{id}_{lang}_{topic_name}/
  step0/
  step1_role/           # Role 库
  step2_themes/         # themes.xlsx + structure_briefs.json（C1/C2/C3）
  C1|C2|C3/
    step3_outlines/
    step4_styles/
    step5_template/
    step6_document/
    step7_pairs/

默认：生成 C1+C2+C3；每档 1 outline × 1 document。
可用 ``--id`` / ``--complexity C1,C2`` 只跑指定档。
默认跑到 step6（成稿）；``--until-step N``（0–7）；默认不跑 step7 配对。
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path

import yaml

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

from LongDocFmtEval.common.draft_resolve import parse_ids_arg  # noqa: E402
from LongDocFmtEval.common.themes_xlsx import COMPLEXITY_TIERS, _norm_complexity, read_themes  # noqa: E402
from LongDocFmtEval.pipeline_auto.auto_paths import (  # noqa: E402
    CONFIG_DIR,
    DATASETS_DIR,
    DEFAULT_CHAPTER_WORKERS,
    DEFAULT_DOCUMENT_COUNT,
    DEFAULT_OUTLINE_COUNT,
    DEFAULT_RUN_THEME_COUNT,
    DEFAULT_THEME_COUNT,
    ROLE_LIBRARY_LLM_CONFIG_PATH,
)
from LongDocFmtEval.pipeline_auto.outline_brief import load_outline_brief  # noqa: E402
from LongDocFmtEval.pipeline_auto.topic_layout import (  # noqa: E402
    document_dir,
    ensure_topic_scaffold,
    outline_dir,
    pairs_dir,
    parse_complexity_arg,
    resolve_topic_dir,
    role_dir,
    sample_id,
    styles_dir,
    template_dir,
    theme_dir,
    themes_dir,
    themes_xlsx_path,
    topic_dir_name,
    write_json,
)

PIPELINE_STEPS = (0, 1, 2, 3, 4, 5, 6, 7)
DEFAULT_UNTIL_STEP = 6


def _normalize_until_step(raw: int | None, *, step7: bool) -> int:
    if raw is None:
        return 7 if step7 else DEFAULT_UNTIL_STEP
    n = int(raw)
    if n not in PIPELINE_STEPS:
        raise ValueError(f"--until-step 须为 {list(PIPELINE_STEPS)} 之一，收到 {n}")
    if step7 and n < 7:
        print(f"[warn] --until-step={n} < 7，忽略 --step7")
    return n


def _pick_field(item: dict, *keys: str) -> str:
    for k in keys:
        v = item.get(k)
        if v is not None and str(v).strip():
            return str(v).strip()
    return ""


def _load_genres() -> list[dict]:
    xlsx = CONFIG_DIR / "genres.xlsx"
    if not xlsx.is_file():
        raise FileNotFoundError(f"缺少 {xlsx}（请先准备体裁表）")
    try:
        from openpyxl import load_workbook  # noqa: PLC0415
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("需要 openpyxl 读取 genres.xlsx") from exc
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h or "").strip().lower() for h in rows[0]]
    out = []
    for raw in rows[1:]:
        if not raw or all(v is None or str(v).strip() == "" for v in raw):
            continue
        item = {headers[i]: raw[i] for i in range(min(len(headers), len(raw)))}
        try:
            gid = int(item.get("id") or item.get("genre_id") or 0)
        except (TypeError, ValueError):
            continue
        if gid <= 0:
            continue
        genre = _pick_field(item, "genre", "genre_key")
        locale = _pick_field(item, "locale") or "zh"
        loc = "en" if locale.lower().startswith("en") else "zh"
        display = _pick_field(
            item,
            "genre_display_name",
            f"genre_display_name_{loc}",
            "display_name",
            "genre_display_name_zh",
            "name",
        ) or genre
        notes = _pick_field(
            item,
            "supplementary_notes",
            f"supplementary_notes_{loc}",
            "notes",
        )
        enabled = str(item.get("enabled") or "true").strip().lower()
        if enabled in {"0", "false", "no", "n"}:
            continue
        status = str(item.get("status") or "").strip().lower()
        if status in {"discarded", "rejected", "disabled"}:
            continue
        out.append({
            "id": gid,
            "genre": genre,
            "locale": loc,
            "display_name": display,
            "notes": notes,
            "domain_id": _pick_field(item, "domain_id"),
            "domain": _pick_field(item, "domain"),
        })
    return out


def _bootstrap_step0(topic: Path, meta: dict, *, overwrite: bool) -> None:
    step0 = topic / "step0"
    step0.mkdir(parents=True, exist_ok=True)
    amplify = {
        "theme_count": meta.get("theme_count"),
        "run_theme_count": meta.get("run_theme_count"),
        "outline_count": meta.get("outline_count"),
        "document_count": meta.get("document_count"),
    }
    if overwrite or not (step0 / "amplify.yaml").is_file():
        (step0 / "amplify.yaml").write_text(
            yaml.safe_dump(amplify, allow_unicode=True, sort_keys=False),
            encoding="utf-8",
        )
    if overwrite or not (step0 / "format_pack.yaml").is_file():
        (step0 / "format_pack.yaml").write_text(
            yaml.safe_dump(
                {
                    "defaults": {
                        "page": {"paper": "A4", "orientation": "portrait"},
                        "body": {
                            "font.ea": "宋体",
                            "size": "12pt",
                            "align": "both",
                            "firstLineIndent": "24pt",
                        },
                    },
                    "presets": {
                        "table_stats": "three_line",
                        "table_cover": "info_underline",
                    },
                    "hints_for_llm": meta.get("notes") or "按体裁惯例区分封面/目录/正文/附录。",
                },
                allow_unicode=True,
                sort_keys=False,
            ),
            encoding="utf-8",
        )
    if overwrite or not (step0 / "constraints.yaml").is_file():
        (step0 / "constraints.yaml").write_text(
            yaml.safe_dump(
                {
                    "exclusive": [["SecCover", "SecFormalHead"]],
                    "require_objects": ["section", "paragraph", "table", "image"],
                    "notes": "自动轨硬约束；step3 选型须遵守。",
                },
                allow_unicode=True,
                sort_keys=False,
            ),
            encoding="utf-8",
        )
    write_json(step0 / "meta.json", meta)


def _load_amplify(
    topic: Path, *, theme_n: int, run_theme_n: int, outline_n: int, document_n: int,
) -> tuple[int, int, int, int]:
    p = topic / "step0" / "amplify.yaml"
    data = {}
    if p.is_file():
        try:
            data = yaml.safe_load(p.read_text(encoding="utf-8")) or {}
        except Exception:  # noqa: BLE001
            data = {}
    theme_count = int(theme_n if theme_n > 0 else data.get("theme_count") or DEFAULT_THEME_COUNT)
    run_theme_count = int(
        run_theme_n if run_theme_n > 0
        else data.get("run_theme_count") or DEFAULT_RUN_THEME_COUNT
    )
    run_theme_count = max(1, min(run_theme_count, theme_count))
    return (
        theme_count,
        run_theme_count,
        int(outline_n if outline_n > 0 else data.get("outline_count") or DEFAULT_OUTLINE_COUNT),
        int(document_n if document_n > 0 else data.get("document_count") or DEFAULT_DOCUMENT_COUNT),
    )


def _sync_role_into_topic(draft_dir: Path, topic: Path) -> Path:
    """Copy Role library into datasets/.../step1_role/ (shared across C1–C3).

    No-op when step1 already wrote directly into the topic role dir.
    """
    dest = role_dir(topic)
    dest.mkdir(parents=True, exist_ok=True)
    try:
        if draft_dir.resolve() == dest.resolve():
            return dest
    except OSError:
        pass
    for name in (
        "draft.yaml", "tree.md", "meta.json", "raw.txt", "inventory.md",
        "llm_usage.json", "error.json", "raw_recovered.txt",
    ):
        src = draft_dir / name
        dst = dest / name
        if not src.is_file():
            continue
        try:
            if src.resolve() == dst.resolve():
                continue
        except OSError:
            pass
        shutil.copy2(src, dst)
    # Copy usage/ folder when present
    src_usage = draft_dir / "usage"
    dst_usage = dest / "usage"
    if src_usage.is_dir():
        try:
            if src_usage.resolve() != dst_usage.resolve():
                if dst_usage.exists():
                    shutil.rmtree(dst_usage)
                shutil.copytree(src_usage, dst_usage)
        except OSError:
            pass
    return dest


def _aggregate_topic_llm_usage(topic: Path) -> dict:
    """Merge every step's llm_usage.json under a topic → llm_usage_total.json."""
    from LongDocFmtEval.common.llm_usage import (  # noqa: PLC0415
        collect_topic_usage_paths,
        merge_usage_reports,
        print_usage_summary,
        write_usage_report,
    )

    topic = Path(topic)
    paths = collect_topic_usage_paths(topic)
    merged = merge_usage_reports(paths)
    try:
        write_usage_report(
            topic / "llm_usage_total.json",
            merged.get("calls") or [],
            step="topic",
            extra={"sources": merged.get("sources") or []},
        )
        write_usage_report(
            topic / "usage" / "llm_usage_total.json",
            merged.get("calls") or [],
            step="topic",
            extra={"sources": merged.get("sources") or []},
        )
    except Exception as e:  # noqa: BLE001
        print(f"[warn] 写出 llm_usage_total.json 失败: {e}")
    summary = merged.get("summary") or {}
    print_usage_summary(summary, prefix="[topic tokens]")
    return summary


def _step1_max_tokens_from_config() -> int | None:
    path = ROLE_LIBRARY_LLM_CONFIG_PATH
    if not path.is_file():
        print(f"[warn] 缺 {path}，step1 max_tokens 不传参")
        return None
    try:
        cfg = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except Exception as exc:  # noqa: BLE001
        print(f"[warn] 读取 {path.name} 失败: {exc}")
        return None
    params = ((cfg.get("language_model") or {}).get("default_params") or {})
    if "max_tokens" not in params or params.get("max_tokens") in (None, ""):
        return None
    try:
        n = int(params["max_tokens"])
    except (TypeError, ValueError):
        return None
    return n if n > 0 else None


def _run_step1(genre_row: dict, *, overwrite: bool) -> Path:
    from LongDocFmtEval.pipeline_auto.scripts import run_step1  # noqa: PLC0415
    from LongDocFmtEval.pipeline_auto.topic_layout import resolve_topic_dir, role_dir  # noqa: PLC0415

    max_tokens = _step1_max_tokens_from_config()
    print(
        f"[step1-config] {ROLE_LIBRARY_LLM_CONFIG_PATH.name} "
        f"max_tokens={max_tokens if max_tokens is not None else '(omit→中转默认)'}"
    )
    topic = resolve_topic_dir(
        int(genre_row["id"]),
        genre=genre_row["genre"],
        locale=genre_row["locale"],
        display_name=genre_row.get("display_name") or "",
    )
    out = role_dir(topic)
    ok = run_step1.run_one(
        genre_row["genre"],
        genre_row["display_name"],
        genre_row["locale"],
        genre_row.get("notes") or "",
        None,
        max_tokens,
        theme_count=0,
        overwrite=overwrite,
        genre_id=int(genre_row["id"]),
        out_dir=out,
    )
    if not ok:
        err = out / "error.json"
        raise RuntimeError(
            f"step1 失败: id={genre_row['id']}"
            + (f"（详见 {err}，marker=STEP1_FAILED）" if err.is_file() else f"（目录 {out}）")
        )
    if not (out / "draft.yaml").is_file():
        raise FileNotFoundError(f"step1 产物未找到: {out / 'draft.yaml'}")
    return out



def _run_step2_themes(
    topic: Path,
    role_library_dir: Path,
    *,
    theme_count: int,
    overwrite: bool,
) -> Path:
    from LongDocFmtEval.pipeline_auto.scripts import run_step2  # noqa: PLC0415

    out = themes_dir(topic)
    out.mkdir(parents=True, exist_ok=True)
    ok = run_step2.run_one(
        role_library_dir,
        count=theme_count,
        library=None,
        genre_arg="",
        locale_arg="",
        extra_notes="",
        temperature=None,
        max_tokens=8192,
        overwrite=overwrite,
        sim_threshold=0.82,
        enable_sim_filter=False,
        out_dir=out,
    )
    if not ok:
        err = out / "error.json"
        raise RuntimeError(
            f"step2 themes 失败: {out}"
            + (f"（详见 {err}，marker=STEP2_FAILED）" if err.is_file() else "")
        )
    return out


def _read_theme_rows(
    topic: Path,
    *,
    complexities: list[str] | None = None,
) -> list[dict]:
    from LongDocFmtEval.pipeline_auto.seed.structure_brief import (  # noqa: PLC0415
        normalize_brief_for_complexity,
        read_structure_briefs,
        validate_complexity_brief_order,
    )

    xlsx = themes_xlsx_path(topic)
    if not xlsx.is_file():
        raise FileNotFoundError(f"缺 themes.xlsx: {xlsx}")
    rows = read_themes(xlsx)
    briefs_file = read_structure_briefs(topic)
    for r in rows:
        cx = _norm_complexity(r.get("complexity"))
        if not cx:
            tid = int(r.get("theme_id") or 0)
            cx = COMPLEXITY_TIERS[tid - 1] if 1 <= tid <= 3 else ""
        r["complexity"] = cx
        brief = briefs_file.get(cx) if cx else None
        if not isinstance(brief, dict):
            brief = {}
        r["structure_brief"] = normalize_brief_for_complexity(
            brief, cx or "C2",
            theme=r.get("theme") or "",
            writing=r.get("writing") or "",
            genre=r.get("genre") or "",
            locale=r.get("locale") or "zh",
        )
    if complexities:
        want = set(complexities)
        rows = [r for r in rows if r.get("complexity") in want]
        order = {c: i for i, c in enumerate(COMPLEXITY_TIERS)}
        rows.sort(key=lambda r: order.get(r.get("complexity") or "", 9))
    # Hard-check order from structure_briefs.json (full C set when present)
    briefs_all = {}
    for cx, brief in briefs_file.items():
        if cx not in COMPLEXITY_TIERS:
            continue
        # find matching theme row for prose context
        match = next((r for r in read_themes(xlsx) if _norm_complexity(r.get("complexity")) == cx), None)
        briefs_all[cx] = normalize_brief_for_complexity(
            brief if isinstance(brief, dict) else {},
            cx,
            theme=(match or {}).get("theme") or "",
            writing=(match or {}).get("writing") or "",
            genre=(match or {}).get("genre") or "",
            locale=(match or {}).get("locale") or "zh",
        )
    if not briefs_all and rows:
        # no json yet: synthesize from complexity presets so pipeline can continue after regenerate
        for r in rows:
            cx = r.get("complexity")
            if cx in COMPLEXITY_TIERS:
                briefs_all[cx] = r["structure_brief"]
    order_errs = validate_complexity_brief_order(briefs_all)
    if order_errs:
        raise RuntimeError("structure_brief C1/C2/C3 顺序校验失败: " + "; ".join(order_errs))
    return rows


def _constraints_notes(topic: Path) -> str:
    bits = []
    fp = topic / "step0" / "format_pack.yaml"
    cp = topic / "step0" / "constraints.yaml"
    if fp.is_file():
        bits.append("format_pack:\n" + fp.read_text(encoding="utf-8")[:2000])
    if cp.is_file():
        bits.append("constraints:\n" + cp.read_text(encoding="utf-8")[:1500])
    return "\n\n".join(bits) if bits else "无"


def _run_outline_for_complexity(
    *,
    topic: Path,
    role_library_dir: Path,
    genre_row: dict,
    theme_row: dict,
    complexity: str,
    overwrite: bool,
) -> Path | None:
    """Generate exactly one outline+stylesheet under ``{Cx}/outline``."""
    from LongDocFmtEval.pipeline_auto.scripts import run_step3  # noqa: PLC0415

    cx_root = theme_dir(topic, complexity, complexity=complexity)
    cx_root.mkdir(parents=True, exist_ok=True)
    write_json(cx_root / "brief.json", {
        "complexity": complexity,
        "theme_id": theme_row.get("theme_id"),
        "theme": theme_row.get("theme"),
        "writing_requirements": theme_row.get("writing") or theme_row.get("writing_requirements"),
        "notes": theme_row.get("notes"),
    })

    base_notes = (
        f"complexity={complexity}\n"
        + (theme_row.get("notes") or "无")
        + "\n\n"
        + _constraints_notes(topic)
    )
    od = outline_dir(topic, complexity=complexity)
    od.mkdir(parents=True, exist_ok=True)
    theme_brief = theme_row.get("structure_brief")
    if not isinstance(theme_brief, dict):
        theme_brief = {}
    ok = run_step3.run_one(
        genre_row["genre"],
        genre_row["locale"],
        str(theme_row.get("theme") or ""),
        str(theme_row.get("writing") or theme_row.get("writing_requirements") or ""),
        base_notes,
        str(role_library_dir / "draft.yaml"),
        None,
        65536,
        step1_dir=role_library_dir,
        overwrite=overwrite,
        theme_key=complexity.lower(),
        genre_id=int(genre_row["id"]),
        out_dir=od,
        prior_outline_briefs="",
        structure_brief=theme_brief,
        complexity=complexity,
    )
    if not ok:
        err_path = od / "error.json"
        detail = ""
        if err_path.is_file():
            try:
                detail = err_path.read_text(encoding="utf-8")[:500]
            except OSError:
                detail = ""
        print(f"[fail] outline {complexity} → 中止该档后续步骤")
        if detail:
            print(f"       {detail[:200]}")
        return None
    brief = load_outline_brief(od)
    if brief:
        brief["outline_id"] = f"{complexity}/outline"
        write_json(od / "outline_brief.json", brief)
    return od


def _simple_collision_gate(template_out: Path) -> dict:
    """Same signature as step4 ``find_collisions`` (section headers + full table props/cells)."""
    rp = template_out / "role_params.json"
    if not rp.is_file():
        return {"ok": False, "reason": "missing role_params.json"}
    try:
        data = json.loads(rp.read_text(encoding="utf-8"))
    except Exception as exc:  # noqa: BLE001
        return {"ok": False, "reason": str(exc)}
    from LongDocFmtEval.pipeline_auto.scripts.run_step4 import find_collisions  # noqa: PLC0415

    groups = find_collisions(data if isinstance(data, dict) else {})
    collisions = []
    for g in groups:
        if len(g) < 2:
            continue
        obj = str((data.get(g[0]) or {}).get("object") or "?")
        for b in g[1:]:
            collisions.append({"object": obj, "a": g[0], "b": b})
    ok = len(collisions) == 0
    report = {"ok": ok, "collisions": collisions}
    write_json(template_out / "validation.json", report)
    return report


def _run_step4_styles(outline: Path, styles_out: Path, *, overwrite: bool) -> bool:
    from LongDocFmtEval.pipeline_auto.scripts import run_step4  # noqa: PLC0415

    ok = run_step4.run_one(
        outline,
        None,
        20480,
        False,
        preview=False,
        overwrite=overwrite,
        figures=False,
        rerender=False,
        out_dir=styles_out,
    )
    if not ok:
        return False
    if overwrite:
        return True
    gate = _simple_collision_gate(styles_out)
    if not gate.get("ok"):
        print(f"[gate] collision: {gate.get('collisions')}")
        write_json(styles_out / "validation.json", {**gate, "policy": "warn_continue_v1"})
    return True


def _run_step5_template(styles_out: Path, tpl_out: Path, *, overwrite: bool, figures: bool,
                        rerender: bool = False, respecimen: bool = False,
                        chapter_workers: int | None = None) -> bool:
    from LongDocFmtEval.pipeline_auto.auto_paths import DEFAULT_CHAPTER_WORKERS  # noqa: PLC0415
    from LongDocFmtEval.pipeline_auto.scripts import run_step5  # noqa: PLC0415

    cw = DEFAULT_CHAPTER_WORKERS if chapter_workers is None else max(1, int(chapter_workers))
    ok = run_step5.run_one(
        styles_out,
        out_dir=tpl_out,
        overwrite=overwrite,
        figures=figures,
        rerender=rerender,
        respecimen=respecimen,
        chapter_workers=cw,
    )
    if not ok:
        return False
    template = tpl_out / "template.docx"
    leftover = tpl_out / "preview.docx"
    if not template.is_file() and leftover.is_file():
        leftover.replace(template)
    elif leftover.is_file() and template.is_file() and leftover.resolve() != template.resolve():
        leftover.unlink()
    return True


def _rerender_topic(
    genre_row: dict, *, complexities: list[str] | None, figures: bool,
) -> dict:
    topic_id = int(genre_row["id"])
    topic = resolve_topic_dir(
        topic_id,
        genre=genre_row.get("genre") or "",
        locale=genre_row.get("locale") or "zh",
        display_name=genre_row.get("display_name") or "",
    )
    tiers = complexities or list(COMPLEXITY_TIERS)
    stats = {"templates": 0, "ok": 0}
    print(f"[rerender] topic={topic.name} complexities={tiers}")
    for cx in tiers:
        styles = styles_dir(topic, complexity=cx)
        tpl = template_dir(topic, complexity=cx)
        if not (styles / "role_params.json").is_file() and not (styles / "template.json").is_file() \
                and not (tpl / "role_params.json").is_file() and not (tpl / "template.json").is_file():
            print(f"  [skip] 缺 role_params.json：{styles.relative_to(topic)}")
            continue
        stats["templates"] += 1
        print(f"\n-- rerender {tpl.relative_to(topic)} --")
        if _run_step5_template(styles, tpl, overwrite=False, figures=figures, rerender=True):
            stats["ok"] += 1
            doc = document_dir(topic, complexity=cx)
            if (doc / "content.json").is_file():
                print(f"-- rerender document {doc.relative_to(topic)} --")
                if _run_step6_document(tpl, doc, overwrite=False, figures=figures, rerender=True):
                    stats["documents"] = stats.get("documents", 0) + 1
                else:
                    print(f"  [warn] step6 重渲失败：{doc}")
            else:
                print(f"  [skip] 无 content.json，不重渲 document.docx：{doc.relative_to(topic)}")
    print(f"[rerender done] {stats['ok']}/{stats['templates']}")
    return {"topic": str(topic), "stats": stats, "mode": "rerender"}


def _run_step6_document(tpl_out: Path, doc_out: Path, *, overwrite: bool, figures: bool,
                        rerender: bool = False, chapter_workers: int | None = None) -> bool:
    from LongDocFmtEval.pipeline_auto.auto_paths import DEFAULT_CHAPTER_WORKERS  # noqa: PLC0415
    from LongDocFmtEval.pipeline_auto.scripts import run_step6  # noqa: PLC0415

    cw = DEFAULT_CHAPTER_WORKERS if chapter_workers is None else max(1, int(chapter_workers))
    return bool(
        run_step6.run_one(
            tpl_out,
            out_dir=doc_out,
            overwrite=overwrite,
            resume=False,
            temperature=0.7,
            max_tokens=16384,
            no_render=False,
            figures=figures,
            rerender=rerender,
            chapter_workers=cw,
        )
    )


def _run_step7_pairs(
    doc_out: Path,
    pair_out: Path,
    *,
    overwrite: bool,
    factors_path: Path | None = None,
) -> bool:
    from LongDocFmtEval.pipeline_auto.seed.build_step7_instance import (  # noqa: PLC0415
        build_step7_instance,
        load_factors_for_step6,
        parse_sample_key_from_step6,
    )
    from LongDocFmtEval.pipeline_auto.seed.seed_ir import load_json  # noqa: PLC0415
    from LongDocFmtEval.pipeline_auto.seed.training_factors import (  # noqa: PLC0415
        DEFAULT_OUTPUT,
        assignments_for_content,
    )

    if not (doc_out / "document.docx").is_file():
        print(f"[skip] step7 缺 document.docx: {doc_out}")
        return False

    from LongDocFmtEval.pipeline_auto.auto_paths import drop_redundant_copies  # noqa: PLC0415
    drop_redundant_copies(doc_out)

    path = factors_path or DEFAULT_OUTPUT
    try:
        tid, cx = parse_sample_key_from_step6(doc_out)
        rows = assignments_for_content(load_json(path), tid, cx)
    except (FileNotFoundError, KeyError, ValueError) as exc:
        print(f"[skip] step7 读 factors 失败: {exc}")
        return False
    if not rows:
        print(f"[skip] step7 无 datasets_factors 行: {tid}_{cx}")
        return False

    ok_modes = 0
    for row in rows:
        im = str(row.get("input_mode") or "").strip().lower()
        if not im:
            continue
        mode_out = pair_out / im
        cached = (mode_out / "sample_info.json").is_file()
        if mode_out.exists() and cached and not overwrite:
            print(f"[cache] step7 已有：{mode_out}")
            ok_modes += 1
            continue
        mode_out.mkdir(parents=True, exist_ok=True)
        try:
            factors = load_factors_for_step6(
                doc_out, global_factors_path=path, input_mode=im,
            )
        except (FileNotFoundError, KeyError) as exc:
            print(f"[skip] step7 无 datasets_factors.json 行 ({im}): {exc}")
            continue
        build_step7_instance(doc_out, mode_out, factors, seed=42)
        ok_modes += 1
    return ok_modes == len(rows) and ok_modes > 0


def run_topic(
    genre_row: dict,
    *,
    theme_count: int,
    complexities: list[str] | None,
    overwrite: bool,
    figures: bool,
    skip_step1: bool,
    until_step: int = DEFAULT_UNTIL_STEP,
    continue_on_error: bool = True,
    chapter_workers: int | None = None,
) -> dict:
    """Build one topic under datasets/{id}_{lang}_{topic_name} with C1/C2/C3 (1 outline × 1 document each)."""
    topic_id = int(genre_row["id"])
    # Hard policy: one outline + one document per complexity
    outline_count = 1
    document_count = 1
    theme_count = max(3, int(theme_count or DEFAULT_THEME_COUNT))

    print(
        f"\n######## TOPIC {topic_id} {genre_row['genre']} → "
        f"{topic_dir_name(topic_id, display_name=genre_row.get('display_name') or '', locale=genre_row['locale'], genre=genre_row['genre'])} ########"
    )
    print(f"[until] 跑到 step{until_step} 为止（含）")
    topic = resolve_topic_dir(
        topic_id,
        genre=genre_row["genre"],
        locale=genre_row["locale"],
        display_name=genre_row.get("display_name") or "",
    )

    ensure_topic_scaffold(topic)
    run_cx = complexities or list(COMPLEXITY_TIERS)
    meta = {
        "id": topic_id,
        "genre": genre_row["genre"],
        "locale": genre_row["locale"],
        "display_name": genre_row["display_name"],
        "notes": genre_row.get("notes") or "",
        "domain_id": genre_row.get("domain_id") or "",
        "theme_count": theme_count,
        "complexities": run_cx,
        "outline_count": outline_count,
        "document_count": document_count,
        "until_step": until_step,
        "track": "pipeline_auto",
        "layout": "datasets/{id}_{lang}_{topic_name}/step1_role+step2_themes+C*/step3..7",
        "no_seed": True,
    }
    _bootstrap_step0(topic, meta, overwrite=overwrite and not skip_step1)
    if until_step <= 0:
        write_json(topic / "step0" / "meta.json", meta)
        write_json(topic / "pipeline_report.json", {"meta": meta, "stats": {}, "stopped_at": 0})
        print(f"[stop] --until-step=0 → 仅 step0：{topic / 'step0'}")
        return {"topic": str(topic), "stats": {}, "stopped_at": 0}

    write_json(topic / "step0" / "meta.json", meta)

    # —— step1 role（共享）——
    role_path = role_dir(topic)
    if skip_step1 and (role_path / "draft.yaml").is_file():
        print(f"[step1] reuse {role_path}")
    else:
        draft = _run_step1(genre_row, overwrite=overwrite)
        role_path = _sync_role_into_topic(draft, topic)
        print(f"[step1] → {role_path}")

    if until_step <= 1:
        stats = {"themes": 0, "outlines": 0, "templates": 0, "documents": 0, "pairs": 0}
        llm_usage = _aggregate_topic_llm_usage(topic)
        write_json(
            topic / "pipeline_report.json",
            {"meta": meta, "stats": stats, "stopped_at": 1, "llm_usage": llm_usage},
        )
        print(f"[stop] --until-step=1 → 请检查 {role_path}")
        return {"topic": str(topic), "stats": stats, "stopped_at": 1}

    # —— step2 themes（整表 C1+C2+C3；运行时可只选部分档）——
    need_themes = (overwrite and not skip_step1) or not themes_xlsx_path(topic).is_file()
    if need_themes:
        _run_step2_themes(
            topic, role_path, theme_count=theme_count, overwrite=overwrite,
        )
    else:
        print(f"[step2] reuse {themes_xlsx_path(topic)}")
    print(f"[step2] → {themes_xlsx_path(topic)}")

    if until_step <= 2:
        stats = {"themes": theme_count, "outlines": 0, "templates": 0, "documents": 0, "pairs": 0}
        llm_usage = _aggregate_topic_llm_usage(topic)
        write_json(
            topic / "pipeline_report.json",
            {"meta": meta, "stats": stats, "stopped_at": 2, "llm_usage": llm_usage},
        )
        print(f"[stop] --until-step=2 → 请检查 {themes_xlsx_path(topic)}")
        return {"topic": str(topic), "stats": stats, "stopped_at": 2}

    theme_rows = _read_theme_rows(topic, complexities=run_cx)
    if not theme_rows:
        raise RuntimeError(
            f"themes.xlsx 无 complexity∈{run_cx} 的主题: {themes_xlsx_path(topic)}"
        )
    print(f"[run] complexities={[r.get('complexity') for r in theme_rows]} (1 outline × 1 document each)")

    stats = {
        "themes": 0, "outlines": 0, "templates": 0, "documents": 0, "pairs": 0,
        "errors": [],
    }
    abort = False
    for trow in theme_rows:
        cx = _norm_complexity(trow.get("complexity")) or ""
        if cx not in COMPLEXITY_TIERS:
            print(f"[skip] 无效 complexity: {trow.get('complexity')}")
            continue
        stats["themes"] += 1
        print(f"\n==== {cx}: {str(trow.get('theme') or '')[:48]} ====")
        od = _run_outline_for_complexity(
            topic=topic,
            role_library_dir=role_path,
            genre_row=genre_row,
            theme_row=trow,
            complexity=cx,
            overwrite=overwrite,
        )
        if od is None:
            stats["errors"].append({
                "complexity": cx,
                "step": 3,
                "reason": "outline 失败或无效（见 C*/outline/error.json）",
            })
            if not continue_on_error:
                print(f"[abort] step3 失败，停止本 topic（可用 --continue-on-error 继续其它档）")
                abort = True
                break
            continue
        if until_step <= 3:
            stats["outlines"] += 1
            continue

        stats["outlines"] += 1
        st = styles_dir(topic, complexity=cx)
        print(f"\n-- step4 {st.relative_to(topic)} --")
        if not _run_step4_styles(od, st, overwrite=overwrite):
            print("[fail] step4")
            stats["errors"].append({"complexity": cx, "step": 4, "reason": "styles 失败"})
            if not continue_on_error:
                print("[abort] step4 失败，停止本 topic")
                abort = True
                break
            continue
        stats["templates"] += 1
        if until_step <= 4:
            continue

        tpl = template_dir(topic, complexity=cx)
        print(f"\n-- step5 {tpl.relative_to(topic)} --")
        if not _run_step5_template(
            st, tpl, overwrite=overwrite, figures=figures, chapter_workers=chapter_workers,
        ):
            print("[fail] step5")
            stats["errors"].append({"complexity": cx, "step": 5, "reason": "template 失败"})
            if not continue_on_error:
                print("[abort] step5 失败，停止本 topic")
                abort = True
                break
            continue
        if until_step <= 5:
            continue

        doc = document_dir(topic, complexity=cx)
        sid = sample_id(topic_id, cx)
        print(f"\n-- step6 {sid} --")
        if not _run_step6_document(
            tpl, doc, overwrite=overwrite, figures=figures, chapter_workers=chapter_workers,
        ):
            print("[fail] step6")
            stats["errors"].append({"complexity": cx, "step": 6, "reason": "document 失败"})
            if not continue_on_error:
                print("[abort] step6 失败，停止本 topic")
                abort = True
                break
            continue
        write_json(doc / "sample_id.json", {"sample_id": sid, "complexity": cx})
        stats["documents"] += 1
        if until_step < 7:
            continue
        pair = pairs_dir(topic, complexity=cx)
        print(f"-- step7 {pair.relative_to(topic)} --")
        if _run_step7_pairs(doc, pair, overwrite=overwrite):
            stats["pairs"] += 1
        else:
            stats["errors"].append({"complexity": cx, "step": 7, "reason": "pairs 失败"})
            if not continue_on_error:
                abort = True
                break

    stopped = until_step
    if abort:
        stopped = -1
        print(f"[abort] topic 中止；errors={stats['errors']}")
    elif until_step <= 3:
        print(f"[stop] --until-step=3 → 请检查 {topic}/C*/step3_outlines/")
    elif until_step <= 4:
        print(f"[stop] --until-step=4 → 请检查 {topic}/C*/step4_styles/")
    elif until_step <= 5:
        print(f"[stop] --until-step=5 → 请检查 {topic}/C*/step5_template/template.docx")
    elif until_step <= 6:
        print(
            "[hint] 已跳过 step7 配对。审改 C*/step5_template 与 C*/step6_document 后，"
            "加 --until-step 7 或 --step7。"
        )
    write_json(
        topic / "pipeline_report.json",
        {
            "meta": meta,
            "stats": stats,
            "stopped_at": stopped,
            "aborted": abort,
            "errors": stats.get("errors") or [],
            "llm_usage": _aggregate_topic_llm_usage(topic),
        },
    )
    if stats.get("errors"):
        write_json(topic / "errors.json", {"errors": stats["errors"], "aborted": abort})

    remediation = None
    try:
        from LongDocFmtEval.pipeline_auto.seed.pipeline_remediation import (  # noqa: PLC0415
            write_topic_remediation,
        )
        remediation = write_topic_remediation(
            topic,
            complexities=tuple(run_cx) if run_cx else None,
        )
        counts = remediation.get("counts") or {}
        print(
            f"[remediation] → {DATASETS_DIR / '_batch_logs' / 'summary.json'} "
            f"manual={counts.get('manual_role_params', 0)} "
            f"retry={counts.get('retry', 0)} "
            f"delete={counts.get('delete_and_retry', 0)}"
        )
    except Exception as exc:  # noqa: BLE001
        print(f"[warn] remediation 报告失败: {exc}")

    print(f"\n[topic done] {topic.name} stats={stats} stopped_at={stopped} aborted={abort}")
    return {
        "topic": str(topic),
        "stats": stats,
        "stopped_at": stopped,
        "aborted": abort,
        "remediation": remediation,
    }


def main() -> None:
    ap = argparse.ArgumentParser(
        description=(
            "pipeline_auto：datasets/{id}_{lang}_{topic_name}/step1_role / step2_themes / C*/step3..7；"
            "role 共享；每档 1 outline × 1 document"
        )
    )
    ap.add_argument("--all", action="store_true", help="对 genres.xlsx 全部 enabled 体裁执行")
    ap.add_argument("--id", dest="topic_id", type=int, default=None, help="指定 topic/体裁 id")
    ap.add_argument("--ids", default="", help="多个 topic id，如 1,26 或 1-3")
    ap.add_argument(
        "--complexity", "--c",
        dest="complexity",
        default="",
        help="只跑指定复杂度档，如 C2 或 C1,C3（默认 C1,C2,C3）",
    )
    ap.add_argument(
        "--themes", type=int, default=DEFAULT_THEME_COUNT,
        help=f"step2 写入 themes.xlsx 条数（默认 {DEFAULT_THEME_COUNT}=C1+C2+C3）",
    )
    ap.add_argument("--overwrite", action="store_true", help="强制重跑已有产物")
    ap.add_argument(
        "--rerender",
        action="store_true",
        help="只重渲已有 C*/step5_template/template.docx 与（若有 content.json）step6_document/document.docx；不重填 T*",
    )
    ap.add_argument("--no-figures", action="store_true", help="step5/6 不出图")
    ap.add_argument(
        "--skip-step1", action="store_true",
        help="保留已有 step1 角色库；若 themes.xlsx 已在，亦保留 step2。配合 --overwrite 从 step3 大纲起重跑",
    )
    ap.add_argument(
        "--until-step",
        type=int,
        default=None,
        metavar="N",
        choices=list(PIPELINE_STEPS),
        help=(
            "跑到第 N 步（含）后停止。0–7。"
            f"默认 {DEFAULT_UNTIL_STEP}（成稿）；检查 template 用 5；配对用 7"
        ),
    )
    ap.add_argument(
        "--step7",
        action="store_true",
        help="未指定 --until-step 时跑到 step7（配对）",
    )
    ap.add_argument(
        "--step6",
        action="store_true",
        help=argparse.SUPPRESS,
    )
    ap.add_argument(
        "--continue-on-error",
        action=argparse.BooleanOptionalAction,
        default=True,
        help="某档 step3/4/5/6 失败时继续其它 complexity（默认开启；--no-continue-on-error 则中止本 topic）",
    )
    ap.add_argument(
        "--chapter-workers", type=int, default=None,
        help=f"step5 specimen / step6 分章写作 LLM 并发（默认 {DEFAULT_CHAPTER_WORKERS}）",
    )
    # Kept for CLI compat but ignored (always 1)
    ap.add_argument("--run-themes", type=int, default=DEFAULT_RUN_THEME_COUNT, help=argparse.SUPPRESS)
    ap.add_argument("--outlines", type=int, default=DEFAULT_OUTLINE_COUNT, help=argparse.SUPPRESS)
    ap.add_argument("--documents", type=int, default=DEFAULT_DOCUMENT_COUNT, help=argparse.SUPPRESS)
    ap.add_argument(
        "--scan-remediation",
        action="store_true",
        help="不跑生成：扫描未完成样本，写出 datasets/_batch_logs/{id}_{Cx}.json 与 summary.json",
    )
    args = ap.parse_args()

    try:
        until_step = _normalize_until_step(args.until_step, step7=bool(args.step7))
        if args.step6 and args.until_step is None and not args.step7:
            print(
                "[warn] --step6 已废弃（旧含义=配对）。"
                "现默认跑到 step6 成稿；配对请用 --until-step 7 或 --step7。"
            )
        complexities = parse_complexity_arg(args.complexity)
    except ValueError as exc:
        ap.error(str(exc))

    genres = _load_genres()
    by_id = {int(g["id"]): g for g in genres}
    if args.all:
        selected = genres
    else:
        ids = parse_ids_arg(args.ids) if args.ids else (
            [args.topic_id] if args.topic_id is not None else []
        )
        if not ids:
            ap.error("请指定 --all 或 --id / --ids")
        missing = [i for i in ids if i not in by_id]
        if missing:
            ap.error(f"genres.xlsx 中无 id={missing}")
        selected = [by_id[i] for i in ids]

    cx_label = ",".join(complexities) if complexities else "C1,C2,C3"
    if args.scan_remediation:
        print(
            f"[scan-remediation] topics={[g['id'] for g in selected]} "
            f"complexity={cx_label}"
        )
    else:
        print(
            f"[pipeline_auto] topics={[g['id'] for g in selected]} "
            f"complexity={cx_label} themes={args.themes} "
            f"until_step={until_step} layout=datasets/{{id}}_{{lang}}_{{topic_name}}"
            + (" mode=rerender" if args.rerender else "")
        )
    print(f"[datasets] {DATASETS_DIR.resolve()}")

    if args.rerender and args.overwrite:
        print("[warn] 同时指定 --rerender 与 --overwrite：本次只按 --rerender")

    from LongDocFmtEval.pipeline_auto.seed.pipeline_remediation import (  # noqa: PLC0415
        scan_topics_and_write,
        write_batch_remediation,
    )

    ok = 0
    if args.scan_remediation:
        topic_paths: list[Path] = []
        for g in selected:
            topic = resolve_topic_dir(
                int(g["id"]),
                genre=str(g.get("genre") or ""),
                locale=str(g.get("locale") or "zh"),
                display_name=str(g.get("display_name") or g.get("topic_name") or ""),
            )
            if not topic.is_dir():
                print(f"[skip] topic 目录不存在: {topic}")
                continue
            topic_paths.append(topic)
        rem = scan_topics_and_write(
            topic_paths,
            complexities=tuple(complexities) if complexities else None,
            label=f"ids={[g['id'] for g in selected]}",
        )
        counts = rem.get("counts") or {}
        print(
            f"[remediation] → {DATASETS_DIR / '_batch_logs' / 'summary.json'} "
            f"manual={counts.get('manual_role_params', 0)} "
            f"retry={counts.get('retry', 0)} "
            f"delete={counts.get('delete_and_retry', 0)}"
        )
        print(f"\n完成扫描 {len(topic_paths)}/{len(selected)} 个 topic")
        return

    for g in selected:
        try:
            if args.rerender:
                _rerender_topic(g, complexities=complexities, figures=not args.no_figures)
                ok += 1
            else:
                run_topic(
                    g,
                    theme_count=int(args.themes),
                    complexities=complexities,
                    overwrite=bool(args.overwrite),
                    figures=not args.no_figures,
                    skip_step1=bool(args.skip_step1),
                    until_step=until_step,
                    continue_on_error=bool(args.continue_on_error),
                    chapter_workers=(
                        int(args.chapter_workers)
                        if args.chapter_workers is not None
                        else DEFAULT_CHAPTER_WORKERS
                    ),
                )
                ok += 1
        except Exception as exc:  # noqa: BLE001
            print(f"[FAIL] topic {g.get('id')}: {exc}")

    batch_path = write_batch_remediation(
        out_dir=DATASETS_DIR / "_batch_logs",
        label=f"ids={[g['id'] for g in selected]}",
    )
    print(f"[remediation] → {batch_path}")

    print(f"\n完成 {ok}/{len(selected)} 个 topic")


# if __name__ == "__main__":
#     main()

if __name__ == "__main__":
    import sys
    # 模拟命令行参数 - 在这里修改 ID 列表
    sys.argv = [
        "run_pipeline_auto.py",

        "--ids", "",
        "--until-step", "6",

    ]
    main()