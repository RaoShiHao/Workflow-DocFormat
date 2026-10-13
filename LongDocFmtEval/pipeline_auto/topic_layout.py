"""Topic-centric layout under ``pipeline_auto/datasets/``.

Layout::

    datasets/{id}_{lang}_{topic_name}/
      step0/
      step1_role/           # Role library (draft.yaml, …)  全 topic 一份
      step2_themes/         # themes.xlsx + structure_briefs.json（C1/C2/C3）
      C1/
        step3_outlines/     # plan.json / stylesheet / selected_library / blocks/
        step4_styles/       # role_params.json（格式）+ template.json（名称目录）
        step5_template/     # specimen + template.docx / assets_preview
        step6_document/     # content.json / document.docx / assets
        step7_pairs/        # 变体；golden 指向 step6 document.docx
      C2/ …
      C3/ …
"""
from __future__ import annotations

import json
import re
import shutil
from pathlib import Path
from typing import Any

from LongDocFmtEval.common.themes_xlsx import COMPLEXITY_TIERS, _norm_complexity
from LongDocFmtEval.pipeline_auto.auto_paths import CONFIG_DIR, DATASETS_DIR, TOPICS_DIR


_FS_UNSAFE = re.compile(r'[<>:"/\\|?*\x00-\x1f]+')
_FS_SPACE = re.compile(r"[\s]+")
_FS_MULTI_UNDERSCORE = re.compile(r"_+")


def _norm_lang(locale: str) -> str:
    return "en" if str(locale or "").lower().startswith("en") else "zh"


def topic_name_slug(topic_name: str, *, max_len: int = 48) -> str:
    """Filesystem-safe slug from genres.xlsx topic display name (keeps CJK)."""
    s = str(topic_name or "").strip()
    s = _FS_UNSAFE.sub("", s)
    s = _FS_SPACE.sub("_", s)
    s = re.sub(r"[^\w\u4e00-\u9fff\-]+", "_", s, flags=re.UNICODE)
    s = _FS_MULTI_UNDERSCORE.sub("_", s).strip("._-")
    if not s:
        s = "topic"
    return s[:max_len]


def topic_dir_name(
    topic_id: int,
    topic_name: str = "",
    locale: str = "zh",
    *,
    genre: str = "",
    display_name: str = "",
) -> str:
    """Canonical folder: ``{id}_{lang}_{topic_name}``."""
    lang = _norm_lang(locale)
    raw = display_name or topic_name or genre or "topic"
    return f"{int(topic_id)}_{lang}_{topic_name_slug(raw)}"


def find_topic_dir_by_id(topic_id: int) -> Path | None:
    """Find existing ``datasets/{id}_*`` (new or legacy ``{id}_topic``)."""
    if not DATASETS_DIR.is_dir():
        return None
    prefix = f"{int(topic_id)}_"
    hits = sorted(
        p for p in DATASETS_DIR.iterdir()
        if p.is_dir() and p.name.startswith(prefix)
    )
    if not hits:
        return None
    modern = [p for p in hits if not p.name.endswith("_topic")]
    return (modern[0] if modern else hits[0]).resolve()


def lookup_genre_row(topic_id: int) -> dict[str, str]:
    """Read id/locale/display_name/genre from config/genres.xlsx (best-effort)."""
    xlsx = CONFIG_DIR / "genres.xlsx"
    empty = {"id": str(int(topic_id)), "locale": "zh", "display_name": "", "genre": ""}
    if not xlsx.is_file():
        return empty
    try:
        from openpyxl import load_workbook  # noqa: PLC0415
    except ImportError:
        return empty
    try:
        wb = load_workbook(xlsx, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception:  # noqa: BLE001
        return empty
    if not rows:
        return empty
    headers = [str(h or "").strip().lower() for h in rows[0]]
    for raw in rows[1:]:
        if not raw:
            continue
        item = {headers[i]: raw[i] for i in range(min(len(headers), len(raw)))}
        try:
            gid = int(item.get("id") or item.get("genre_id") or 0)
        except (TypeError, ValueError):
            continue
        if gid != int(topic_id):
            continue
        locale = str(item.get("locale") or "zh").strip() or "zh"
        loc = _norm_lang(locale)
        display = (
            str(item.get("genre_display_name") or "").strip()
            or str(item.get(f"genre_display_name_{loc}") or "").strip()
            or str(item.get("display_name") or "").strip()
            or str(item.get("genre_display_name_zh") or "").strip()
            or str(item.get("name") or "").strip()
        )
        genre = str(item.get("genre") or item.get("genre_key") or "").strip()
        return {
            "id": str(gid),
            "locale": loc,
            "display_name": display or genre,
            "genre": genre,
        }
    return empty


def theme_name(theme_id: int) -> str:
    """Legacy helper; prefer complexity folders C1/C2/C3."""
    return f"theme_{int(theme_id):02d}"


def outline_name(outline_id: int = 1) -> str:
    return f"outline_{int(outline_id):02d}"


def document_name(document_id: int = 1) -> str:
    return f"document_{int(document_id):02d}"


def resolve_topic_dir(
    topic_id: int,
    *,
    genre: str = "",
    locale: str = "zh",
    topic_name: str = "",
    display_name: str = "",
) -> Path:
    """Return ``datasets/{id}_{lang}_{topic_name}/`` (create if missing).

    Reuses an existing ``{id}_*`` folder when present (including legacy ``{id}_topic``).
    If only a legacy folder exists and a real topic name is known, rename to the canonical form.
    """
    DATASETS_DIR.mkdir(parents=True, exist_ok=True)
    info = lookup_genre_row(topic_id)
    # Caller overrides win; otherwise fill from genres.xlsx
    disp = display_name or topic_name or info.get("display_name") or ""
    gen = genre or info.get("genre") or ""
    if genre or display_name or topic_name:
        loc = _norm_lang(locale)
    else:
        loc = _norm_lang(info.get("locale") or locale)

    canonical = DATASETS_DIR / topic_dir_name(
        topic_id, topic_name=disp or gen, locale=loc, genre=gen, display_name=disp,
    )
    existing = find_topic_dir_by_id(topic_id)
    if existing is not None:
        if existing.resolve() == canonical.resolve():
            return existing
        legacy_like = existing.name.endswith("_topic") or existing.name != canonical.name
        if legacy_like and not canonical.exists():
            try:
                existing.rename(canonical)
                print(f"[topic] rename {existing.name} → {canonical.name}")
                return canonical.resolve()
            except OSError as exc:
                print(f"[topic] rename failed ({exc}); reuse {existing.name}")
                return existing
        return existing

    canonical.mkdir(parents=True, exist_ok=True)
    return canonical.resolve()


def find_legacy_topic_dir(topic_id: int) -> Path | None:
    """Optional: locate old ``work/topics/topic_XX_*`` if present."""
    if not TOPICS_DIR.is_dir():
        return None
    prefix = f"topic_{int(topic_id):02d}_"
    hits = [p for p in TOPICS_DIR.iterdir() if p.is_dir() and p.name.startswith(prefix)]
    return hits[0].resolve() if hits else None


def ensure_topic_scaffold(topic: Path) -> dict[str, Path]:
    paths = {
        "step0": topic / "step0",
        "step1_role": topic / "step1_role",
        "step2_themes": topic / "step2_themes",
        "C1": topic / "C1",
        "C2": topic / "C2",
        "C3": topic / "C3",
    }
    for p in paths.values():
        p.mkdir(parents=True, exist_ok=True)
    for cx in COMPLEXITY_TIERS:
        for sub in ("step3_outlines", "step4_styles", "step5_template", "step6_document", "step7_pairs"):
            (topic / cx / sub).mkdir(parents=True, exist_ok=True)
    return paths


def role_dir(topic: Path) -> Path:
    """Shared Role library: ``step1_role/`` (legacy: ``role/``, ``step1/``)."""
    modern = topic / "step1_role"
    for cand in (modern, topic / "role", topic / "step1"):
        if (cand / "draft.yaml").is_file():
            return cand
    for cand in (topic / "role", topic / "step1"):
        if cand.is_dir():
            try:
                next(cand.iterdir())
            except StopIteration:
                continue
            else:
                return cand
    modern.mkdir(parents=True, exist_ok=True)
    return modern


def themes_dir(topic: Path) -> Path:
    """Step2 themes directory: ``step2_themes/`` (themes.xlsx + structure_briefs.json)."""
    modern = topic / "step2_themes"
    if (modern / "themes.xlsx").is_file():
        return modern
    # legacy: themes.xlsx at topic root / step2 /
    if (topic / "themes.xlsx").is_file():
        return topic
    if (topic / "step2" / "themes.xlsx").is_file():
        return topic / "step2"
    modern.mkdir(parents=True, exist_ok=True)
    return modern


def themes_xlsx_path(topic: Path) -> Path:
    d = themes_dir(topic)
    p = d / "themes.xlsx"
    if p.is_file():
        return p
    for alt in (
        topic / "themes.xlsx",
        topic / "step2_themes" / "themes.xlsx",
        topic / "step2" / "themes.xlsx",
        topic / "step1_role" / "themes.xlsx",
        topic / "role" / "themes.xlsx",
        topic / "step1" / "themes.xlsx",
    ):
        if alt.is_file():
            return alt
    return topic / "step2_themes" / "themes.xlsx"


def structure_briefs_json_path(topic: Path) -> Path:
    """Preferred ``step2_themes/structure_briefs.json`` with legacy fallbacks."""
    modern = themes_dir(topic) / "structure_briefs.json"
    if modern.is_file():
        return modern
    for alt in (
        topic / "structure_briefs.json",
        topic / "step2_themes" / "structure_briefs.json",
        topic / "step2" / "structure_briefs.json",
    ):
        if alt.is_file():
            return alt
    return themes_dir(topic) / "structure_briefs.json"


def complexity_dir(topic: Path, complexity: str) -> Path:
    cx = _norm_complexity(complexity) or str(complexity or "").strip().upper()
    if cx not in COMPLEXITY_TIERS:
        raise ValueError(f"complexity 须为 C1/C2/C3，收到 {complexity!r}")
    return topic / cx


def theme_dir(topic: Path, theme_id: int | str, *, complexity: str | None = None) -> Path:
    """Per-complexity workspace root (``C1`` / ``C2`` / ``C3``)."""
    if complexity:
        return complexity_dir(topic, complexity)
    if isinstance(theme_id, str) and _norm_complexity(theme_id) in COMPLEXITY_TIERS:
        return complexity_dir(topic, theme_id)
    try:
        tid = int(theme_id)
    except (TypeError, ValueError):
        tid = 0
    if 1 <= tid <= 3:
        return complexity_dir(topic, COMPLEXITY_TIERS[tid - 1])
    legacy = topic / "step3" / theme_name(int(tid or 1))
    return legacy if legacy.is_dir() else complexity_dir(topic, "C1")


def outline_dir(
    topic: Path,
    theme_id: int | str = 1,
    outline_id: int = 1,
    *,
    complexity: str | None = None,
) -> Path:
    """``{Cx}/step3_outlines`` (legacy: ``outline/``)."""
    del outline_id
    root = theme_dir(topic, theme_id, complexity=complexity)
    modern = root / "step3_outlines"
    if (modern / "plan.json").is_file() or (modern.is_dir() and any(modern.glob("*"))):
        return modern
    legacy = root / "outline"
    if (legacy / "plan.json").is_file():
        return legacy
    if complexity is None:
        try:
            tid = int(theme_id) if not isinstance(theme_id, str) or theme_id.isdigit() else 1
        except (TypeError, ValueError):
            tid = 1
        for base in ("step3", "step2"):
            old = topic / base / theme_name(tid) / outline_name(1)
            if (old / "plan.json").is_file():
                return old
    modern.mkdir(parents=True, exist_ok=True)
    return modern


def styles_dir(
    topic: Path,
    theme_id: int | str = 1,
    outline_id: int = 1,
    *,
    complexity: str | None = None,
) -> Path:
    """``{Cx}/step4_styles`` (legacy combined: ``step4_template``)."""
    del outline_id
    root = theme_dir(topic, theme_id, complexity=complexity)
    modern = root / "step4_styles"
    if (modern / "role_params.json").is_file() or (modern / "template.json").is_file():
        return modern
    combined = root / "step4_template"
    split_tpl = root / "step5_template"
    split_ready = (split_tpl / "template.docx").is_file() or (split_tpl / "specimen_plan.json").is_file()
    if (
        ((combined / "template.json").is_file() or (combined / "role_params.json").is_file())
        and not split_ready
    ):
        return combined
    modern.mkdir(parents=True, exist_ok=True)
    return modern


def template_dir(
    topic: Path,
    theme_id: int | str = 1,
    outline_id: int = 1,
    *,
    complexity: str | None = None,
) -> Path:
    """``{Cx}/step5_template`` (legacy: ``step4_template`` / ``template/``)."""
    del outline_id
    root = theme_dir(topic, theme_id, complexity=complexity)
    modern = root / "step5_template"
    if (
        (modern / "template.docx").is_file()
        or (modern / "specimen_plan.json").is_file()
        or (modern.is_dir() and any(modern.glob("*")))
    ):
        return modern
    combined = root / "step4_template"
    if (combined / "template.docx").is_file() or (combined / "template.json").is_file():
        return combined
    legacy = root / "template"
    if (legacy / "template.json").is_file():
        return legacy
    if complexity is None:
        try:
            tid = int(theme_id) if not isinstance(theme_id, str) or str(theme_id).isdigit() else 1
        except (TypeError, ValueError):
            tid = 1
        for base in ("step4", "step3"):
            old = topic / base / theme_name(tid) / outline_name(1)
            if (old / "template.json").is_file():
                return old
    modern.mkdir(parents=True, exist_ok=True)
    return modern


def document_dir(
    topic: Path,
    theme_id: int | str = 1,
    outline_id: int = 1,
    document_id: int = 1,
    *,
    complexity: str | None = None,
) -> Path:
    """``{Cx}/step6_document`` (legacy: ``step5_document`` / ``document/``)."""
    del outline_id, document_id
    root = theme_dir(topic, theme_id, complexity=complexity)
    modern = root / "step6_document"
    if (
        (modern / "document.docx").is_file()
        or (modern / "content.json").is_file()
        or (modern.is_dir() and any(modern.glob("*")))
    ):
        return modern
    legacy_named = root / "step5_document"
    if (
        (legacy_named / "document.docx").is_file()
        or (legacy_named / "content.json").is_file()
        or (legacy_named.is_dir() and any(legacy_named.glob("*")))
    ):
        return legacy_named
    legacy = root / "document"
    if (legacy / "document.docx").is_file() or (legacy / "content.json").is_file():
        return legacy
    if complexity is None:
        try:
            tid = int(theme_id) if not isinstance(theme_id, str) or str(theme_id).isdigit() else 1
        except (TypeError, ValueError):
            tid = 1
        old = topic / "step5" / theme_name(tid) / outline_name(1) / document_name(1)
        if old.is_dir():
            return old
    modern.mkdir(parents=True, exist_ok=True)
    return modern


def pairs_dir(
    topic: Path,
    theme_id: int | str = 1,
    outline_id: int = 1,
    document_id: int = 1,
    *,
    complexity: str | None = None,
) -> Path:
    """``{Cx}/step7_pairs`` (legacy: ``step6_pairs`` / ``pairs/``)."""
    del outline_id, document_id
    root = theme_dir(topic, theme_id, complexity=complexity)
    modern = root / "step7_pairs"
    for cand in (modern, root / "step6_pairs", root / "pairs"):
        if cand.is_dir() and any(cand.iterdir()):
            return cand
    modern.mkdir(parents=True, exist_ok=True)
    return modern


def _rewrite_path_text(raw: str, mapping: list[tuple[str, str]]) -> str:
    out = raw
    for old, new in mapping:
        out = out.replace(old, new)
        out = out.replace(old.replace("\\", "/"), new.replace("\\", "/"))
        out = out.replace(old.replace("/", "\\"), new.replace("/", "\\"))
    return out


def migrate_complexity_step_dirs(cx_root: Path) -> list[str]:
    """Split ``step4_template`` and rename step5/6 folders. Idempotent."""
    notes: list[str] = []
    root = Path(cx_root)
    if not root.is_dir():
        return notes

    style_names = {
        "role_params.json", "template.json", "template.md", "style_layers.json",
        "validation.json", "raw.txt", "source.json", "brief.json", "llm_usage.json",
    }
    style_dirs = {"usage"}
    combined = root / "step4_template"
    styles = root / "step4_styles"
    template = root / "step5_template"
    if combined.is_dir() and not (styles / "template.json").is_file() and not (styles / "role_params.json").is_file():
        styles.mkdir(parents=True, exist_ok=True)
        template.mkdir(parents=True, exist_ok=True)
        for p in list(combined.iterdir()):
            name = p.name
            dest_root = styles if (name in style_names or name in style_dirs) else template
            dest = dest_root / name
            if dest.exists():
                notes.append(f"skip exists {dest.relative_to(root)}")
                continue
            shutil.move(str(p), str(dest))
            notes.append(f"{combined.name}/{name} → {dest_root.name}/{name}")
        try:
            next(combined.iterdir())
        except StopIteration:
            combined.rmdir()
            notes.append(f"removed empty {combined.name}")

    mapping = [
        ("step4_template", "step5_template"),
        ("step5_document", "step6_document"),
        ("step6_pairs", "step7_pairs"),
    ]
    skip_rewrite = {
        "role_params.json", "template.json", "content.json", "plan.json",
        "stylesheet.json", "selected_library.yaml",
    }
    for old_name, new_name in (
        ("step5_document", "step6_document"),
        ("step6_pairs", "step7_pairs"),
    ):
        src, dest = root / old_name, root / new_name
        if src.is_dir() and not dest.exists():
            src.rename(dest)
            notes.append(f"{old_name} → {new_name}")

    for folder in (styles, template, root / "step6_document", root / "step7_pairs"):
        if not folder.is_dir():
            continue
        folder_map = (
            [("step4_template", "step4_styles"), ("step5_document", "step6_document"),
             ("step6_pairs", "step7_pairs")]
            if folder == styles else mapping
        )
        for p in folder.rglob("*.json"):
            if p.name in skip_rewrite:
                continue
            try:
                text = p.read_text(encoding="utf-8")
            except OSError:
                continue
            new = _rewrite_path_text(text, folder_map)
            if p.name in ("source.json", "source_step5.json", "source_step6.json", "source_step7.json"):
                new = new.replace("step5_template/role_params.json", "step4_styles/role_params.json")
                new = new.replace("step5_template/template.json", "step4_styles/template.json")
            if new != text:
                p.write_text(new, encoding="utf-8")
                notes.append(f"rewrote paths in {p.relative_to(root).as_posix()}")

    doc = root / "step6_document"
    if doc.is_dir():
        meta = {
            "step": 6,
            "template_dir": str((root / "step5_template").resolve()).replace("\\", "/"),
            "styles_dir": str((root / "step4_styles").resolve()).replace("\\", "/"),
        }
        old = doc / "source_step5.json"
        newp = doc / "source_step6.json"
        if old.is_file() and not newp.is_file():
            old.replace(newp)
            notes.append("source_step5.json → source_step6.json")
        if (doc / "document.docx").is_file() or (doc / "content.json").is_file() or newp.is_file():
            try:
                data = json.loads(newp.read_text(encoding="utf-8")) if newp.is_file() else {}
            except json.JSONDecodeError:
                data = {}
            data.update(meta)
            newp.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    pairs = root / "step7_pairs"
    if pairs.is_dir():
        old = pairs / "source_step6.json"
        newp = pairs / "source_step7.json"
        if old.is_file() and not newp.is_file():
            old.replace(newp)
            notes.append("source_step6.json → source_step7.json")
        if newp.is_file():
            try:
                data = json.loads(newp.read_text(encoding="utf-8"))
            except json.JSONDecodeError:
                data = {}
            data["step"] = 7
            data["document_dir"] = str((root / "step6_document").resolve()).replace("\\", "/")
            data["step6_dir"] = data["document_dir"]
            if "step5_dir" in data:
                data["step5_dir"] = data["document_dir"]
            newp.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    tpl = root / "step5_template"
    if tpl.is_dir():
        (tpl / "source_step5.json").write_text(
            json.dumps({
                "step": 5,
                "styles_dir": str((root / "step4_styles").resolve()).replace("\\", "/"),
                "template_dir": str(tpl.resolve()).replace("\\", "/"),
            }, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    return notes


def migrate_topic_step_dirs(topic: Path) -> list[str]:
    notes: list[str] = []
    topic = Path(topic)
    for cx in COMPLEXITY_TIERS:
        n = migrate_complexity_step_dirs(topic / cx)
        notes.extend(f"{cx}: {msg}" for msg in n)
    return notes


def sample_id(topic_id: int, complexity: str, **_kwargs) -> str:
    cx = _norm_complexity(complexity) or str(complexity)
    return f"{int(topic_id)}__{cx}"


def parse_complexity_arg(text: str | None) -> list[str] | None:
    """Parse ``C1``, ``C1,C2``, ``c2`` → list of tiers; None/empty → all three."""
    if text is None or not str(text).strip():
        return None
    out: list[str] = []
    for part in re.split(r"[,，\s]+", str(text).strip()):
        if not part:
            continue
        cx = _norm_complexity(part)
        if not cx:
            raise ValueError(f"无效 complexity: {part!r}（仅 C1/C2/C3）")
        if cx not in out:
            out.append(cx)
    return out or None


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))
