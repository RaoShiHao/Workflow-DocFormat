"""Shared helpers for the per-genre themes Excel (**step2** → consumed by step3+).

Columns: theme_id, complexity, genre, locale, theme, writing_requirements,
supplementary_notes, enabled.

Stored under ``step2_themes/themes.xlsx``. C1/C2/C3 budgets live in sibling
``structure_briefs.json`` (not this xlsx).
"""
from __future__ import annotations

from pathlib import Path

THEME_COLUMNS = [
    "theme_id",
    "complexity",
    "genre",
    "locale",
    "theme",
    "writing_requirements",
    "supplementary_notes",
    "enabled",
]

COMPLEXITY_TIERS = ("C1", "C2", "C3")


def _norm_complexity(raw) -> str:
    s = str(raw or "").strip().upper()
    if s in COMPLEXITY_TIERS:
        return s
    for t in COMPLEXITY_TIERS:
        if s == t[1:] or s.endswith(t) or s == f"TIER{t[1:]}" or s == f"LEVEL{t[1:]}":
            return t
    return ""


def _assign_theme_ids(rows: list[dict]) -> list[dict]:
    used: set[int] = set()
    out: list[dict] = []
    for r in rows:
        rr = dict(r)
        tid = rr.get("theme_id", "")
        try:
            n = int(tid) if str(tid).strip() not in ("", "None", "nan") else None
        except (TypeError, ValueError):
            n = None
        if n is not None and n > 0 and n not in used:
            rr["theme_id"] = n
            used.add(n)
        else:
            rr["theme_id"] = None
        out.append(rr)
    nxt = 1
    for rr in out:
        if rr["theme_id"] is None:
            while nxt in used:
                nxt += 1
            rr["theme_id"] = nxt
            used.add(nxt)
            nxt += 1
    return out


def make_themes_template(path: Path, blank_rows: int = 8, rows: list[dict] | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "themes"
    ws.append(THEME_COLUMNS)
    prepared = _assign_theme_ids(list(rows or []))
    for r in prepared:
        if not r.get("complexity"):
            tid = int(r.get("theme_id") or 0)
            r["complexity"] = COMPLEXITY_TIERS[tid - 1] if 1 <= tid <= 3 else ""
        ws.append([r.get(c, "") for c in THEME_COLUMNS])
    blank = {"locale": "zh", "enabled": 1, "complexity": ""}
    start_id = (max((int(r["theme_id"]) for r in prepared), default=0) + 1) if prepared else 1
    for i in range(blank_rows):
        row = {**blank, "theme_id": start_id + i}
        ws.append([row.get(c, "") for c in THEME_COLUMNS])

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="4F81BD")
        cell.font = Font(bold=True, color="FFFFFF")
    ws.freeze_panes = "A2"
    widths = {
        "theme_id": 10,
        "complexity": 12,
        "genre": 18,
        "locale": 8,
        "theme": 36,
        "writing_requirements": 48,
        "supplementary_notes": 28,
        "enabled": 9,
    }
    wrap = {"theme", "writing_requirements", "supplementary_notes"}
    for idx, col in enumerate(THEME_COLUMNS):
        ws.column_dimensions[get_column_letter(idx + 1)].width = widths.get(col, 16)
        if col in wrap:
            for row in ws.iter_rows(min_row=2):
                row[idx].alignment = Alignment(wrap_text=True, vertical="top")
    en_col = get_column_letter(THEME_COLUMNS.index("enabled") + 1)
    dv = DataValidation(type="list", formula1='"1,0"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{en_col}2:{en_col}{ws.max_row}")
    cx_col = get_column_letter(THEME_COLUMNS.index("complexity") + 1)
    dv2 = DataValidation(type="list", formula1='"C1,C2,C3"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add(f"{cx_col}2:{cx_col}{ws.max_row}")

    ws2 = wb.create_sheet("说明")
    for line in [
        ["字段", "含义"],
        ["theme_id", "本库内主题编号（从 1 起）"],
        ["complexity", "C1/C2/C3 相对格式复杂度"],
        ["theme", "选题（宽泛）"],
        ["writing_requirements", "写作结构约束（相对该档）"],
        ["enabled", "1=参与"],
        ["", ""],
        ["结构预算", "见同目录 structure_briefs.json（按 C1/C2/C3；不在本表）"],
        ["来源", "step2 生成 themes.xlsx + structure_briefs.json"],
    ]:
        ws2.append(line)
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 92

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"[themes] {path}")


def read_themes(path: Path, *, enabled_only: bool = True) -> list[dict]:
    import pandas as pd

    df = pd.read_excel(path, sheet_name=0).fillna("")
    raw_rows: list[dict] = []
    for i, r in enumerate(df.iterrows(), 1):
        _, row = r
        genre = str(row.get("genre", "") or "").strip()
        theme = str(row.get("theme", "") or "").strip()
        if not genre or not theme:
            continue
        enabled = str(row.get("enabled", 1)).strip() not in ("0", "false", "False")
        tid_raw = row.get("theme_id", "")
        try:
            tid = int(tid_raw) if str(tid_raw).strip() not in ("", "None", "nan") else i
        except (TypeError, ValueError):
            tid = i
        cx = _norm_complexity(row.get("complexity", ""))
        if not cx and 1 <= tid <= 3:
            cx = COMPLEXITY_TIERS[tid - 1]
        raw_rows.append({
            "theme_id": tid,
            "complexity": cx,
            "genre": genre,
            "locale": str(row.get("locale", "") or "zh").strip() or "zh",
            "theme": theme,
            "writing": str(row.get("writing_requirements", "") or "").strip(),
            "notes": str(row.get("supplementary_notes", "") or "").strip(),
            "enabled": enabled,
        })
    prepared = _assign_theme_ids(raw_rows)
    if enabled_only:
        return [r for r in prepared if r.get("enabled", True)]
    return prepared


def ensure_theme_ids_file(path: Path) -> int:
    rows = read_themes(path, enabled_only=False)
    excel_rows = []
    for r in rows:
        excel_rows.append({
            "theme_id": r["theme_id"],
            "complexity": r.get("complexity") or "",
            "genre": r["genre"],
            "locale": r["locale"],
            "theme": r["theme"],
            "writing_requirements": r["writing"],
            "supplementary_notes": r["notes"],
            "enabled": 1 if r.get("enabled", True) else 0,
        })
    make_themes_template(path, blank_rows=0, rows=excel_rows)
    return len(excel_rows)
